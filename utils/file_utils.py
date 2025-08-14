import os
from io import BytesIO
import PyPDF2
import docx
from bs4 import BeautifulSoup
import openpyxl

def extract_text_from_file(file_path: str, file_bytes: bytes = None) -> str:
    """
    Extrae texto de archivos PDF, DOCX, HTML, TXT y XLSX.

    Args:
        file_path (str): Ruta del archivo.
        file_bytes (bytes, optional): Bytes del archivo. Si no se provee, se lee del disco.

    Returns:
        str: Texto extraído del archivo.
    """
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if file_bytes is None:
            with open(file_path, 'rb') as f:
                file_bytes = f.read()
        if ext == '.pdf':
            reader = PyPDF2.PdfReader(BytesIO(file_bytes))
            for page in reader.pages:
                text += page.extract_text() or ""
        elif ext == '.docx':
            doc = docx.Document(BytesIO(file_bytes))
            text = "\n".join([para.text for para in doc.paragraphs])
        elif ext in ['.html', '.htm']:
            soup = BeautifulSoup(file_bytes, 'html.parser')
            text = soup.get_text(separator='\n', strip=True)
        elif ext == '.txt':
            text = file_bytes.decode('utf-8', errors='ignore')
        elif ext == '.xlsx':
            wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    text += '\t'.join([str(cell) if cell is not None else '' for cell in row]) + '\n'
    except Exception as e:
        print(f"Error extrayendo texto del archivo {file_path}: {e}")
    return text

def save_uploaded_file(upload_dir: str, filename: str, file_bytes: bytes) -> str:
    """Guarda un archivo subido en el directorio especificado."""
    os.makedirs(upload_dir, exist_ok=True)
    file_path = os.path.join(upload_dir, filename)
    with open(file_path, 'wb') as f:
        f.write(file_bytes)
    return file_path

def allowed_file_extension(filename: str, allowed_extensions: set = None) -> bool:
    """Valida si la extensión del archivo es permitida."""
    if allowed_extensions is None:
        allowed_extensions = {'.pdf', '.docx', '.html', '.htm', '.txt', '.xlsx'}
    ext = os.path.splitext(filename)[1].lower()
    return ext in allowed_extensions
